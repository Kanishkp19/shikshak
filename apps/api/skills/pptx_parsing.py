"""
Shikshak AI — PPTX parsing skill.

Wraps python-pptx with a single function returning slide-level text.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any


def parse_pptx(file_path: str | Path) -> dict[str, Any]:
    """Parse a PPTX file into slide-level text chunks.

    Returns:
        {
            "page_count": int,
            "pages": [{"page": int, "section_label": str, "text": str}]
        }
    """
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "python-pptx not installed — pip install python-pptx"
        ) from e

    pres = Presentation(str(file_path))
    pages: list[dict] = []
    for i, slide in enumerate(pres.slides, start=1):
        chunks: list[str] = []
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = (shape.text_frame.text or "").replace("\x00", "").replace("\u0000", "").strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    title = text
                chunks.append(text)
        pages.append(
            {
                "page": i,
                "section_label": f"Slide {i}: {title}" if title else f"Slide {i}",
                "text": "\n".join(chunks),
            }
        )
    return {"page_count": len(pages), "pages": pages}
